"""Master Sending/Pick-Up schedule store.

`XlsxSheetStore` is the local backend (openpyxl); a `GoogleSheetStore` can be added
later with the same `upsert(direction, row) -> "inserted"|"updated"` contract.
Direction selects the tab, so within a tab the row key is the Name (= Name + Direction).
"""
from pathlib import Path

import openpyxl

COLUMNS = [
    "Name", "Position", "Airlines / Flight",
    "Date", "Time", "Dispatch Time", "Airport", "Terminal", "Notes", "Update History",
]
_TABS = {"sendoff": "Sending", "pickup": "Pick-Up"}


class XlsxSheetStore:
    def __init__(self, path):
        self.path = Path(path)
        if self.path.exists():
            self.wb = openpyxl.load_workbook(self.path)
        else:
            self.wb = openpyxl.Workbook()
            first = self.wb.active
            first.title = "Pick-Up"
            first.append(COLUMNS)
            self.wb.create_sheet("Sending").append(COLUMNS)
            self.wb.save(self.path)

    def upsert(self, direction: str, row: dict) -> str:
        """Insert or update the row for row['Name'] on the direction's tab."""
        ws = self.wb[_TABS[direction]]
        target = next((r for r in ws.iter_rows(min_row=2) if r[0].value == row["Name"]), None)
        if target is None:
            ws.append([row.get(col, "") for col in COLUMNS])
            result = "inserted"
        else:
            for col, cell in zip(COLUMNS, target):
                if col in row:
                    cell.value = row[col]
            result = "updated"
        self.wb.save(self.path)
        return result
