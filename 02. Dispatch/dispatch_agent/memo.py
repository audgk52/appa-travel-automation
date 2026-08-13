"""Travel Memo readers.

`MemoSource` is the swap point: `DocxMemoSource` parses a local .docx now; a
`GoogleDocMemoSource` can be added later returning the same `TravelMemo` shape.
"""
import re
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path

import docx

# Korean airports we dispatch to/from (drives pick-up and send-off detection).
KOREA_AIRPORTS = {"ICN", "GMP"}

_FLIGHT_NO = re.compile(r"^[A-Z0-9]{2,3}\s?\d+")  # OZ 1085, UA 805, ANA 862
_IATA = re.compile(r"\b([A-Z]{3})\b")  # airport code, wherever it sits in the cell
_TIME = re.compile(r"(\d{1,2}:\d{2})")
_TERMINAL = re.compile(r"Terminal\s+(\d)")  # numbered only; "TBD"/none -> no terminal
_PLUS_DAYS = re.compile(r"\+\s*(\d+)\s*day")
_NAME_ROLE = re.compile(r"^(.*?)\s*\((.*)\)\s*$")
# A passenger-cell line that is a cast/guest label rather than a person's name.
_LABEL = re.compile(r"Cast|Guest|Crew|Staff|#|\d")


@dataclass
class FlightLeg:
    date: date  # departure date of this leg (see korea_arrival for arrival-date adjustment)
    airline: str
    flight_no: str
    from_code: str
    from_terminal: str | None
    to_code: str
    to_terminal: str | None
    depart: str
    arrive: str
    arrive_offset_days: int = 0

    @property
    def terminal(self) -> str | None:
        """Korea-side terminal for this leg (from side if departing Korea, to side if arriving)."""
        if self.from_code in KOREA_AIRPORTS:
            return self.from_terminal
        if self.to_code in KOREA_AIRPORTS:
            return self.to_terminal
        return None


@dataclass
class TravelMemo:
    passenger: str  # primary traveler
    role: str  # primary traveler's role/position ("" if none)
    legs: list
    travelers: list | None = None  # every traveler on the memo (defaults to [passenger])
    group_note: str = ""  # role + cast/guest labels, for 특이사항 pre-fill

    def __post_init__(self):
        if self.travelers is None:
            self.travelers = [self.passenger]
        if not self.group_note:
            self.group_note = self.role

    @property
    def korea_departure(self) -> FlightLeg | None:
        return next((leg for leg in self.legs if leg.from_code in KOREA_AIRPORTS), None)

    @property
    def korea_arrival(self) -> FlightLeg | None:
        leg = next((leg for leg in self.legs if leg.to_code in KOREA_AIRPORTS), None)
        if leg is None:
            return None
        return replace(leg, date=leg.date + timedelta(days=leg.arrive_offset_days))


def _parse_date(cell: str) -> date:
    # "Sunday May 17, 2026" / "Monday Jun 22, 2026" -> date(...)
    _, rest = cell.strip().split(" ", 1)
    rest = rest.strip()
    for fmt in ("%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(rest, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognized date: {cell!r}")


def _parse_city(cell: str):
    # Handles both memo formats:
    #   "ICN / South Korea Terminal 1"                        -> ("ICN", "1")
    #   "Incheon Int, Incheon, South Korea  ICN  Terminal 2"  -> ("ICN", "2")
    # The 3-letter IATA code is the first standalone all-caps triple, wherever it sits.
    m = _IATA.search(cell)
    code = m.group(1) if m else cell.strip().split()[0]
    t = _TERMINAL.search(cell)
    return code, (t.group(1) if t else None)


def _parse_passengers(cell: str):
    """Split a passenger cell into (primary_name, primary_role, all_names, group_note).

    A memo may list a family/group across lines, some of which are cast/guest
    labels (e.g. "Cast #2 Guests") rather than names. Names -> 탑승자; the primary
    role plus any labels -> group_note (the 특이사항 pre-fill).
    """
    names, labels, primary_role = [], [], None

    def add_role(role):
        nonlocal primary_role
        if primary_role is None:
            primary_role = role
        else:
            labels.append(role)

    for line in (ln.strip() for ln in cell.splitlines() if ln.strip()):
        m = _NAME_ROLE.match(line)
        if m and m.group(1).strip():  # "Name (Role)"
            names.append(m.group(1).strip())
            add_role(m.group(2).strip())
        elif m:  # "(Role)" on its own line -> role for the preceding traveler
            add_role(m.group(2).strip())
        elif _LABEL.search(line):  # "Cast #2", "Cast #2 Guests"
            labels.append(line)
        else:
            names.append(line)
    if not names:
        names = [cell.strip()]
    primary = names[0]
    role = primary_role or ""
    parts = list(dict.fromkeys(([primary_role] if primary_role else []) + labels))
    return primary, role, names, " & ".join(parts)


def _parse_time(cell: str):
    # "16:05 + 1 day" -> ("16:05", 1)
    m = _TIME.match(cell.strip())
    time = m.group(1) if m else cell.strip()
    off = _PLUS_DAYS.search(cell)
    return time, (int(off.group(1)) if off else 0)


class DocxMemoSource:
    """Parse a Travel Memo from a local .docx file."""

    def load(self, path) -> TravelMemo:
        doc = docx.Document(str(path))
        passenger, role, travelers, group_note = self._passenger(doc)
        legs = self._legs(doc)
        return TravelMemo(
            passenger=passenger,
            role=role,
            legs=legs,
            travelers=travelers,
            group_note=group_note,
        )

    @staticmethod
    def _passenger(doc):
        cell = doc.tables[0].rows[0].cells[1].text.strip()
        return _parse_passengers(cell)

    @staticmethod
    def _legs(doc):
        legs = []
        for row in doc.tables[1].rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            flight_no = " ".join(cells[2].split())
            if not _FLIGHT_NO.match(flight_no):
                continue  # header row or the disclaimer row
            from_code, from_term = _parse_city(cells[3])
            to_code, to_term = _parse_city(cells[4])
            depart, _ = _parse_time(cells[5])
            arrive, arr_off = _parse_time(cells[6])
            legs.append(
                FlightLeg(
                    date=_parse_date(cells[0]),
                    airline=cells[1],
                    flight_no=flight_no,
                    from_code=from_code,
                    from_terminal=from_term,
                    to_code=to_code,
                    to_terminal=to_term,
                    depart=depart,
                    arrive=arrive,
                    arrive_offset_days=arr_off,
                )
            )
        return legs
