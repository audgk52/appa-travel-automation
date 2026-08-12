"""Master schedule store — one tab, one row per dispatch, sorted by Send Date.

`ScheduleStore` is the local (openpyxl) backend; a `GoogleSheetStore` can be added
later with the same `upsert(row) -> "inserted"|"updated"` contract. Row key = Name + Direction.
"""
from datetime import date, datetime
from pathlib import Path

import openpyxl

TAB = "Schedule"
COLUMNS = [
    "Send Date", "Dispatch Date", "Direction", "Name", "Position", "Flight",
    "Airport", "Terminal", "Flight Time", "Dispatch Time", "Notes", "Message", "Rev / Updated",
]
_NAME = COLUMNS.index("Name")
_DIRECTION = COLUMNS.index("Direction")
_SEND_DATE = COLUMNS.index("Send Date")


def _as_date(value):
    """Normalize a cell value to a date for sorting (openpyxl reloads dates as datetime)."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.min


class ScheduleStore:
    def __init__(self, path):
        self.path = Path(path)
        if self.path.exists():
            self.wb = openpyxl.load_workbook(self.path)
        else:
            self.wb = openpyxl.Workbook()
            self.wb.active.title = TAB
            self.wb[TAB].append(COLUMNS)
            self.wb.save(self.path)

    def upsert(self, row: dict) -> str:
        ws = self.wb[TAB]
        key = (row["Name"], row["Direction"])
        target = next(
            (r for r in ws.iter_rows(min_row=2)
             if (r[_NAME].value, r[_DIRECTION].value) == key),
            None,
        )
        if target is None:
            ws.append([row.get(col, "") for col in COLUMNS])
            result = "inserted"
        else:
            for col, cell in zip(COLUMNS, target):
                if col in row:
                    cell.value = row[col]
            result = "updated"
        self._sort_by_send_date(ws)
        self.wb.save(self.path)
        return result

    @staticmethod
    def _sort_by_send_date(ws):
        data = [
            [c.value for c in r]
            for r in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in r)
        ]
        data.sort(key=lambda vals: (vals[_SEND_DATE] is None, _as_date(vals[_SEND_DATE])))
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for vals in data:
            ws.append(vals)
