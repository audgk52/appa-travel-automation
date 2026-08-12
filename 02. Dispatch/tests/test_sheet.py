"""Master Sending/Pick-Up sheet store (PE-3) — one .xlsx, two tabs, upsert by Name."""
import openpyxl

from dispatch_agent.sheet import XlsxSheetStore


def test_creates_file_with_two_tabs(tmp_path):
    XlsxSheetStore(tmp_path / "master.xlsx").upsert(
        "sendoff", {"Name": "Alex Mosley", "Dispatch Time": "10:30"}
    )
    wb = openpyxl.load_workbook(tmp_path / "master.xlsx")
    assert wb.sheetnames == ["Pick-Up", "Sending"]


def test_sendoff_upsert_updates_existing_row_by_name(tmp_path):
    p = tmp_path / "master.xlsx"
    store = XlsxSheetStore(p)
    assert store.upsert("sendoff", {"Name": "Alex Mosley", "Dispatch Time": "10:30"}) == "inserted"
    assert store.upsert("sendoff", {"Name": "Alex Mosley", "Dispatch Time": "11:00"}) == "updated"

    ws = openpyxl.load_workbook(p)["Sending"]
    names = [r[0].value for r in ws.iter_rows(min_row=2) if r[0].value]
    assert names == ["Alex Mosley"]  # one row, not two


def test_header_has_airport_next_to_terminal_and_no_reservation(tmp_path):
    p = tmp_path / "master.xlsx"
    XlsxSheetStore(p).upsert("sendoff", {"Name": "X"})
    header = [c.value for c in openpyxl.load_workbook(p)["Sending"][1]]
    assert "Reservation #" not in header
    assert header.index("Airport") + 1 == header.index("Terminal")


def test_pickup_goes_to_pickup_tab(tmp_path):
    p = tmp_path / "master.xlsx"
    XlsxSheetStore(p).upsert("pickup", {"Name": "Nicole Lang"})
    wb = openpyxl.load_workbook(p)
    assert wb["Pick-Up"].max_row == 2  # header + one row
    assert wb["Sending"].max_row == 1  # header only
