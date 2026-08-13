"""CLI helpers — schedule-row assembly, direction detection, revision color."""
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from dispatch_agent.builder import build_record
from dispatch_agent.cli import build_schedule_row, directions_in_memo, parse_color
from dispatch_agent.memo import DocxMemoSource

REF = Path(__file__).resolve().parent.parent / "Reference Files"
MEMO = REF / "APPA_Travel_Memo_63_Elizabeth Tedder_051526.docx"


def _memo():
    return DocxMemoSource().load(MEMO)


def test_directions_in_memo_returns_pickup_then_sendoff():
    assert directions_in_memo(_memo()) == ["pickup", "sendoff"]


def test_parse_color_from_filename():
    assert parse_color("APPA_Travel_Memo_63_Elizabeth Tedder_BLUE_052926.docx") == "BLUE"
    assert parse_color("APPA_Travel_Memo_63_Elizabeth Tedder_PINK_061726.docx") == "PINK"
    assert parse_color("APPA_Travel_Memo_63_Elizabeth Tedder_051526.docx") == "original"


def test_build_schedule_row_sendoff():
    rec = build_record(_memo(), "sendoff", "VIP")
    row = build_schedule_row(_memo(), rec, "sendoff", color="PINK")
    assert row["Direction"] == "sendoff"
    assert row["Name"] == "Elizabeth Marie Tedder"
    assert row["Dispatch Date"] == date(2026, 6, 22)
    assert row["Send Date"] == date(2026, 6, 20)  # dispatch - 2
    assert row["Flight"] == "AS 120"
    assert row["Airport"] == "ICN"
    assert row["Terminal"] == "1"
    assert row["Flight Time"] == "19:35"
    assert row["Dispatch Time"] == "15:35"
    assert row["Notes"] == "US Line Producer (VIP)"  # role prefill + ad-hoc
    assert "공항 샌딩" in row["Message"]
    assert row["Rev / Updated"].startswith("PINK")


def test_build_schedule_row_pickup():
    rec = build_record(_memo(), "pickup", "")
    row = build_schedule_row(_memo(), rec, "pickup")
    assert row["Direction"] == "pickup"
    assert row["Dispatch Date"] == date(2026, 5, 18)
    assert row["Send Date"] == date(2026, 5, 16)
    assert row["Flight"] == "AC 63"
    assert "공항 픽업" in row["Message"]


from dispatch_agent.cli import sync_row_to_calendar


class _OkCalendar:
    def upsert_event(self, row):
        return "created"


class _FailCalendar:
    def upsert_event(self, row):
        raise RuntimeError("boom")


def test_sync_row_to_calendar_success():
    status, err = sync_row_to_calendar(_OkCalendar(), {"Name": "X", "Direction": "pickup"})
    assert status == "created"
    assert err is None


def test_sync_row_to_calendar_surfaces_failure():
    status, err = sync_row_to_calendar(_FailCalendar(), {"Name": "X", "Direction": "pickup"})
    assert status is None
    assert isinstance(err, RuntimeError)


def test_main_calendar_disabled_when_unconfigured(tmp_path, monkeypatch, capsys):
    monkeypatch.delenv("APPA_GOOGLE_SA_KEY", raising=False)
    monkeypatch.delenv("APPA_GCAL_CALENDAR_ID", raising=False)
    sheet = tmp_path / "master.xlsx"
    from dispatch_agent.cli import main

    main(["--memo", str(MEMO), "--notes", "N/A", "--sheet", str(sheet)])
    out = capsys.readouterr().out
    assert "[calendar] disabled" in out


def test_main_configured_failure_exits_nonzero_but_saves_local(tmp_path, monkeypatch, capsys):
    # Configured but Calendar init fails: WARN to stderr, exit non-zero, yet the
    # local schedule must still be written (local-first invariant).
    monkeypatch.setenv("APPA_GOOGLE_SA_KEY", "/tmp/key.json")
    monkeypatch.setenv("APPA_GCAL_CALENDAR_ID", "cal@x")
    import dispatch_agent.cli as climod

    def boom(_key_path):
        raise RuntimeError("no creds")

    monkeypatch.setattr(climod, "build_calendar_service", boom)
    sheet = tmp_path / "master.xlsx"
    with pytest.raises(SystemExit) as exc:
        climod.main(["--memo", str(MEMO), "--notes", "N/A", "--sheet", str(sheet)])
    assert exc.value.code == 1
    assert "[WARN]" in capsys.readouterr().err
    assert sheet.exists()
    assert openpyxl.load_workbook(sheet)["Schedule"].max_row >= 2  # local rows written


def test_main_partial_config_exits_nonzero_but_saves_local(tmp_path, monkeypatch, capsys):
    # Only one env var set -> CalendarConfigError must not crash before local save.
    monkeypatch.setenv("APPA_GOOGLE_SA_KEY", "/tmp/key.json")
    monkeypatch.delenv("APPA_GCAL_CALENDAR_ID", raising=False)
    from dispatch_agent.cli import main

    sheet = tmp_path / "master.xlsx"
    with pytest.raises(SystemExit) as exc:
        main(["--memo", str(MEMO), "--notes", "N/A", "--sheet", str(sheet)])
    assert exc.value.code == 1
    assert "[WARN]" in capsys.readouterr().err
    assert sheet.exists()
    assert openpyxl.load_workbook(sheet)["Schedule"].max_row >= 2  # local rows written


def test_main_configured_success_prints_calendar_status(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("APPA_GOOGLE_SA_KEY", "/tmp/key.json")
    monkeypatch.setenv("APPA_GCAL_CALENDAR_ID", "cal@x")
    import dispatch_agent.cli as climod

    class _FakeCalendar:
        def __init__(self, *a, **k):
            pass

        def upsert_event(self, row):
            return "created"

    monkeypatch.setattr(climod, "build_calendar_service", lambda _p: object())
    monkeypatch.setattr(climod, "CalendarSync", _FakeCalendar)
    sheet = tmp_path / "master.xlsx"
    climod.main(["--memo", str(MEMO), "--notes", "N/A", "--sheet", str(sheet)])
    assert "calendar: created" in capsys.readouterr().out
