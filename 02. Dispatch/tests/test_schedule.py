"""Single-tab master schedule store (Phase A) — one row per dispatch, sorted by Send Date."""
from datetime import date

import openpyxl

from dispatch_agent.schedule import COLUMNS, ScheduleStore


def _row(name, direction, send, dispatch, message="msg"):
    return {
        "Send Date": send, "Dispatch Date": dispatch, "Direction": direction,
        "Name": name, "Message": message,
    }


def test_creates_single_tab_with_header(tmp_path):
    ScheduleStore(tmp_path / "m.xlsx").upsert(_row("A", "sendoff", date(2026, 6, 20), date(2026, 6, 22)))
    wb = openpyxl.load_workbook(tmp_path / "m.xlsx")
    assert wb.sheetnames == ["Schedule"]
    assert [c.value for c in wb["Schedule"][1]] == COLUMNS


def test_upsert_inserts_then_updates_same_key(tmp_path):
    p = tmp_path / "m.xlsx"
    store = ScheduleStore(p)
    assert store.upsert(_row("Alex Mosley", "sendoff", date(2026, 6, 20), date(2026, 6, 22))) == "inserted"
    assert store.upsert(_row("Alex Mosley", "sendoff", date(2026, 6, 21), date(2026, 6, 23))) == "updated"
    ws = openpyxl.load_workbook(p)["Schedule"]
    names = [r[3].value for r in ws.iter_rows(min_row=2) if r[3].value]
    assert names == ["Alex Mosley"]  # one row, not two


def test_same_name_different_direction_are_two_rows(tmp_path):
    p = tmp_path / "m.xlsx"
    store = ScheduleStore(p)
    store.upsert(_row("Beth", "pickup", date(2026, 5, 16), date(2026, 5, 18)))
    store.upsert(_row("Beth", "sendoff", date(2026, 6, 20), date(2026, 6, 22)))
    ws = openpyxl.load_workbook(p)["Schedule"]
    assert ws.max_row == 3  # header + 2


def test_rows_sorted_ascending_by_send_date(tmp_path):
    p = tmp_path / "m.xlsx"
    store = ScheduleStore(p)
    store.upsert(_row("Late", "sendoff", date(2026, 6, 20), date(2026, 6, 22)))
    store.upsert(_row("Early", "pickup", date(2026, 5, 16), date(2026, 5, 18)))
    ws = openpyxl.load_workbook(p)["Schedule"]
    names_in_order = [r[3].value for r in ws.iter_rows(min_row=2) if r[3].value]
    assert names_in_order == ["Early", "Late"]


def test_upsert_after_reload_mixes_date_and_datetime(tmp_path):
    # openpyxl returns saved dates as datetime; a fresh row is a date. Sort must not crash.
    p = tmp_path / "m.xlsx"
    ScheduleStore(p).upsert(_row("Late", "sendoff", date(2026, 6, 20), date(2026, 6, 22)))
    ScheduleStore(p).upsert(_row("Early", "pickup", date(2026, 5, 16), date(2026, 5, 18)))
    ws = openpyxl.load_workbook(p)["Schedule"]
    names = [r[3].value for r in ws.iter_rows(min_row=2) if r[3].value]
    assert names == ["Early", "Late"]


def test_stores_message_text(tmp_path):
    p = tmp_path / "m.xlsx"
    ScheduleStore(p).upsert(_row("A", "sendoff", date(2026, 6, 20), date(2026, 6, 22), message="▷날짜 : 2026.06.22"))
    ws = openpyxl.load_workbook(p)["Schedule"]
    msg_col = COLUMNS.index("Message")
    assert ws.cell(row=2, column=msg_col + 1).value == "▷날짜 : 2026.06.22"
